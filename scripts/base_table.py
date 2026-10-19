# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "google-api-python-client>=2.100",
#   "google-auth>=2.20",
#   "openpyxl>=3.1",
# ]
# ///
"""
底表：从 schemas/base-table.schema.json 生成并核对底表（Google 表格）。

  uv run scripts/base_table.py xlsx                  生成样例工作簿 templates/base-table-sample.xlsx
  uv run scripts/base_table.py build  --target main   建/补正式底表的结构（不动已有数据行）
  uv run scripts/base_table.py build  --target sample 建样例表，并用假数据覆盖其数据行
  uv run scripts/base_table.py check  --target main   核对表结构与 schema 一致，不一致退出码 1

回流脚本（如 scripts/gsc_pull.py）复用本文件的读写层：google_service、read_table、upsert_rows、
replace_rows、refresh_progress。

配置（都在两个仓库之外）：
  凭据  GOOGLE_APPLICATION_CREDENTIALS，否则 ~/.config/de-linker-playbook/service-account.json
  表格 ID  PLAYBOOK_CONFIG，否则 ~/.config/de-linker-playbook/config.toml：
          [base_table]
          spreadsheet_id = "..."         # 正式底表
          sample_spreadsheet_id = "..."  # 样例表
首页、各「说明」页每次 build 整页重写；数据表只写表头、格式、下拉，从不清数据行。
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import tomllib
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = ROOT / "schemas" / "base-table.schema.json"
XLSX_OUT = ROOT / "templates" / "base-table-sample.xlsx"
CONFIG_DIR = Path.home() / ".config" / "de-linker-playbook"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]
VALIDATION_ROWS = 2000
RETRIES = 5          # googleapiclient retries 429 / 5xx / socket timeouts with backoff
HTTP_TIMEOUT = 120   # seconds per request
BEIJING = timezone(timedelta(hours=8))

# 表头颜色 = 谁写这一列
WRITER_FILL = {"脚本": "DDEBF7", "AI": "E2EFDA", "AI（据口述）": "FCE4D6"}
GREY = "D9D9D9"
SAMPLE = "样例，假数据"

# ---------------------------------------------------------------- content


def load_schema() -> dict:
    return json.loads(SCHEMA.read_text(encoding="utf-8"))


def columns(spec: dict) -> list[str]:
    return list(spec["properties"])


def writer_key(meta: dict) -> str:
    w = meta["x-writer"]
    if w.startswith("AI（据口述）"):
        return "AI（据口述）"
    return "AI" if w.startswith("AI") else "脚本"


# A page is a list of (cells, style). style: title | h2 | header | text | 脚本 | AI | AI（据口述）
Page = list[tuple[list[str], str]]


def home_page(schema: dict, sample: bool) -> Page:
    defs = schema["$defs"]
    if sample:
        head = [(["底表（样例）——这本表里的数据全是假的，只用来看懂表的样子"], "title"),
                (["网址里的 sample、公司名里的「样例公司」、Message-ID 里的 example.invalid 都是编的。正式数据在另一本「底表」里。"], "text")]
    else:
        head = [(["底表"], "title"),
                (["正式数据。数据表由脚本和 AI 写入；想看填满了是什么样，打开另一本「底表（样例）」。"], "text")]
    page: Page = head + [
        ([], "text"),
        (["这本表是干什么的"], "h2"),
        (["记下我们做过的每一件事（发布网页、发邮件……），再把 Google、邮箱、网站询盘传回来的结果一条条记在旁边，"
          "这样就能看出：每次动作走到了哪一步，下一步该做什么。看板（Looker Studio）直接读这本表画图；AI 也读这本表给出下一步。"], "text"),
        ([], "text"),
        (["工作表一览（每张数据表后面紧跟一张同名的「说明」页）"], "h2"),
        (["工作表", "回答什么问题", "一行代表什么"], "header"),
    ]
    for name in schema["x-sheet-order"]:
        page.append(([name, defs[name]["description"], defs[name]["x-row"]], "text"))
    page += [
        ([], "text"),
        (["反馈分四级：一次动作走到了哪一步"], "h2"),
        (["级别", "意思", "网页上看什么", "邮件上看什么"], "header"),
        (["L0 已执行", "我们做了（执行登记里有这一行）", "页面已上线", "邮件已发出，在「已发送」里找到"], "text"),
        (["L1 已到达", "对方那边收到了", "Google 收录了这页，或这页在搜索结果里被显示（有展示）", "没有退信，推断已送达"], "text"),
        (["L2 有互动", "有人回应了", "有人从搜索结果点进来（有点击）", "对方真人回信"], "text"),
        (["L3 有询盘", "有人来要报价了（北极星）", "网站上提交了形式发票申请", "同左"], "text"),
        (["未到 L1", "发了但没到", "—", "邮件被退回（退信）"], "text"),
        ([], "text"),
        (["各表怎么串起来：靠「登记编号」"], "h2"),
        (["执行登记给每次动作一个编号（如 E-20261010-01）。网页收录、网页表现按网址找到这个编号；"
          "邮件反馈按邮件的 Message-ID 找到它；询盘有证据时由 AI 填上它。动作进度把同一个编号的所有消息汇总成一行。"], "text"),
        (["例（见「底表（样例）」）：E-20261012-02 是给样例公司甲发的第 1 封信 → 邮件反馈里它先「已发送里找到」(L0)、"
          "再「未见退信」(L1)、再「真人回复」(L2) → 询盘表里 90001 归因到它 (L3) → 动作进度里这一行的当前最高级别就是 L3。"], "text"),
        ([], "text"),
        (["表头颜色 = 谁写这一列"], "h2"),
        (["蓝色", "脚本自动写（从 Google、邮箱、网站读来）"], "脚本"),
        (["绿色", "AI 写（整理、判断、归因）"], "AI"),
        (["橙色", "你在对话里说一句，AI 查证后替你写"], "AI（据口述）"),
        (["你不需要在任何一张表里打字；看到不对的地方，在对话里告诉 AI。"], "text"),
        ([], "text"),
        (["给初学者的读表小贴士"], "h2"),
        (["1. 第一行是表头，已冻结：往下滚动时表头一直在。"], "text"),
        (["2. 想只看某一类（比如只看「获客」线）：选中表头那一行 → 菜单「数据」→「创建过滤器」→ 点列名旁边的小漏斗勾选。不会改动数据。"], "text"),
        (["3. 有下拉箭头的列只能填固定的几个值，可选的值写在对应的说明页里。"], "text"),
        (["4. 时刻都是北京时间，写成「年-月-日 时:分」。德国比北京晚 6 小时（冬令时晚 7 小时）。网页表现、搜索词的日期按 Google 的美国太平洋时间切日，和北京的「哪一天」可能差一天。"], "text"),
        (["5. 网页表现里某个组合某天没有行，表示 Google 那天没有这个组合的数据，当作 0。搜索词表加起来比网页表现少是正常的（隐私原因）。"], "text"),
        (["6. 不知道某列什么意思：打开同名的说明页，每一列都写了是什么、从哪来、谁写、空着表示什么。"], "text"),
        ([], "text"),
        (["这本表不放什么"], "h2"),
        (["人名、邮箱地址、电话、地址、邮件正文、询盘备注都不放；公司名可以放。表格只共享给所有者账号、日常账号和脚本用的服务账号。"], "text"),
        (["询盘表只含网站上的形式发票申请；联系表单和直接发到邮箱的询盘现在读不到，所以询盘数偏低。"], "text"),
    ]
    return page


def guide_page(name: str, spec: dict) -> Page:
    page: Page = [
        ([f"「{name}」这张表怎么读"], "title"),
        ([], "text"),
        (["这张表回答什么"], "h2"), ([spec["description"]], "text"),
        (["一行代表什么"], "h2"), ([spec["x-row"]], "text"),
        (["什么时候会多出行"], "h2"), ([spec["x-rows-added"]], "text"),
    ]
    if "x-not-counted" in spec:
        page += [(["要知道的缺口"], "h2"), ([spec["x-not-counted"]], "text")]
    page += [
        (["哪几列合起来不重复"], "h2"),
        (["「" + "」+「".join(spec["x-key"]) + "」：同样的组合只会有一行，脚本重跑不会多出重复行。"], "text"),
        ([], "text"),
        (["每一列（第一格的颜色 = 谁写）"], "h2"),
        (["列名", "这列是什么", "从哪来", "谁写", "取值示例", "只能填这些值", "空着表示什么", "怎么读、注意什么"], "header"),
    ]
    for col, meta in spec["properties"].items():
        must = "（必填）" if col in spec.get("required", []) else ""
        page.append(([col + must, meta["description"], meta["x-source"], meta["x-writer"],
                      "；".join(str(x) for x in meta.get("examples", [])), "；".join(meta.get("enum", [])),
                      meta["x-empty"], meta.get("x-read", "")], writer_key(meta)))
    for col, meta in spec["properties"].items():
        if "x-states" not in meta:
            continue
        english: dict[str, list[str]] = {}
        for en, zh in meta.get("x-map", {}).items():
            english.setdefault(zh, []).append(en)
        page += [([], "text"), ([f"「{col}」每种取值是什么意思"], "h2"),
                 (["取值", "属于哪个收录结论", "是什么意思", "Google 的英文原文"], "header")]
        for state, info in meta["x-states"].items():
            page.append(([state, info["收录结论"], info["意思"], "；".join(english.get(state, ["—"]))], "text"))
    return page


HOME_WIDTHS = [16, 48, 52, 40]
GUIDE_WIDTHS = [22, 36, 36, 14, 26, 26, 28, 40]


def px(width: int) -> int:
    return width * 7 + 10


def merge_span(widths: list[int]) -> int:
    """Single-cell text rows are merged across this many columns (about 110 width units)."""
    total = 0
    for i, w in enumerate(widths, start=1):
        total += w
        if total >= 110:
            return i
    return len(widths)


def text_row_px(text: str, style: str, widths: list[int]) -> int:
    char = {"title": 22, "h2": 17}.get(style, 14)
    line = {"title": 30, "h2": 24}.get(style, 21)
    avail = sum(px(w) for w in widths[:merge_span(widths)]) - 12
    used = sum(char if ord(ch) > 255 else char * 0.55 for ch in text)
    return max(1, -(-int(used) // avail)) * line + 8

# ---------------------------------------------------------------- sample data

# Raw sample rows use the API's own values (coverageState, pageFetchState, country codes …) under
# keys starting with "_"; translate_* turns them into the table's Chinese columns the same way the
# pullback scripts will.
SAMPLE_ROWS: dict[str, list[dict]] = {
    "执行登记": [
        {"登记编号": "E-20261010-01", "登记时刻（北京时间）": "2026-10-10 17:30", "执行日期": "2026-10-10",
         "线": "营销", "渠道": "网站", "动作": "新发布页面", "内容说明": "样例：发布一篇英文选型指南",
         "网址": "https://de-linker.com/sample-guide-a/", "对应交付物": "marketing/05-content/sample-guide-a.md",
         "来源": "操作者口述", "查证结果": "已查证", "查证证据": "HTTP 200，页面标题与口述一致",
         "查证时刻（北京时间）": "2026-10-10 17:31"},
        {"登记编号": "E-20261011-01", "登记时刻（北京时间）": "2026-10-11 15:10", "执行日期": "2026-10-11",
         "线": "营销", "渠道": "网站", "动作": "改正文", "内容说明": "样例：给德文产品页补一张参数表",
         "网址": "https://de-linker.com/de/sample-produkt-b/", "来源": "操作者口述", "查证结果": "已查证",
         "查证证据": "HTTP 200，参数表已出现在页面上", "查证时刻（北京时间）": "2026-10-11 15:12"},
        {"登记编号": "E-20261012-01", "登记时刻（北京时间）": "2026-10-12 20:00", "执行日期": "2026-10-12",
         "线": "营销", "渠道": "网站", "动作": "改网址（旧址 301 跳新址）", "内容说明": "样例：案例页换成更短的地址",
         "网址": "https://de-linker.com/sample-new-c/", "旧网址": "https://de-linker.com/sample-old-c/",
         "来源": "新站 PR", "查证结果": "已查证",
         "查证证据": "新地址 HTTP 200；旧地址 301 跳到新地址；证据为新站仓库已合并的 PR（样例）",
         "查证时刻（北京时间）": "2026-10-12 20:02"},
        {"登记编号": "E-20261012-02", "登记时刻（北京时间）": "2026-10-12 23:20", "执行日期": "2026-10-12",
         "线": "获客", "渠道": "邮件", "动作": "发出首封", "内容说明": "样例：给样例公司甲发第 1 封外展邮件",
         "收件公司": "样例公司甲 GmbH", "邮件 Message-ID": "<sample-001@example.invalid>",
         "外展序列": "S-样例公司甲", "第几封": 1, "计划共几封": 3, "下一封计划日期": "2026-10-19",
         "对应交付物": "lead/06-outreach/sample-sequence-a.md", "来源": "操作者口述", "查证结果": "已查证",
         "查证证据": "在「已发送」里找到这封信", "查证时刻（北京时间）": "2026-10-13 14:00"},
        {"登记编号": "E-20261013-01", "登记时刻（北京时间）": "2026-10-13 16:40", "执行日期": "2026-10-13",
         "线": "获客", "渠道": "邮件", "动作": "发出首封", "内容说明": "样例：给样例公司乙发第 1 封外展邮件",
         "收件公司": "样例公司乙 AG", "邮件 Message-ID": "<sample-002@example.invalid>",
         "外展序列": "S-样例公司乙", "第几封": 1, "计划共几封": 3, "下一封计划日期": "2026-10-20",
         "对应交付物": "lead/06-outreach/sample-sequence-b.md", "来源": "操作者口述", "查证结果": "已查证",
         "查证证据": "在「已发送」里找到这封信", "查证时刻（北京时间）": "2026-10-14 14:00"},
        {"登记编号": "E-20261014-01", "登记时刻（北京时间）": "2026-10-14 18:05", "执行日期": "2026-10-14",
         "线": "营销", "渠道": "LinkedIn", "动作": "发帖", "内容说明": "样例：在公司 LinkedIn 发一条产品帖",
         "网址": "https://www.linkedin.com/feed/update/sample-0001/", "来源": "操作者口述",
         "查证结果": "未查证（无法公开访问）", "查证证据": "需要登录才能打开，按操作者口述记录",
         "查证时刻（北京时间）": "2026-10-14 18:06"},
    ],
    "网页收录": [
        {"检查日期": "2026-10-11", "网址": "https://de-linker.com/sample-guide-a/",
         "_coverage": "URL is unknown to Google"},
        {"检查日期": "2026-10-12", "网址": "https://de-linker.com/sample-guide-a/",
         "_coverage": "Submitted and indexed", "_robots": "ALLOWED", "_indexing": "INDEXING_ALLOWED",
         "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-12 06:05",
         "Google 选的正式地址": "https://de-linker.com/sample-guide-a/",
         "页面声明的正式地址": "https://de-linker.com/sample-guide-a/"},
        {"检查日期": "2026-10-14", "网址": "https://de-linker.com/sample-guide-a/",
         "_coverage": "Submitted and indexed", "_robots": "ALLOWED", "_indexing": "INDEXING_ALLOWED",
         "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-13 21:40",
         "Google 选的正式地址": "https://de-linker.com/sample-guide-a/",
         "页面声明的正式地址": "https://de-linker.com/sample-guide-a/"},
        {"检查日期": "2026-10-14", "网址": "https://de-linker.com/de/sample-produkt-b/",
         "_coverage": "Crawled - currently not indexed", "_robots": "ALLOWED", "_indexing": "INDEXING_ALLOWED",
         "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-12 11:40",
         "Google 选的正式地址": "https://de-linker.com/de/sample-produkt-b/",
         "页面声明的正式地址": "https://de-linker.com/de/sample-produkt-b/"},
        {"检查日期": "2026-10-14", "网址": "https://de-linker.com/sample-new-c/",
         "_coverage": "Discovered - currently not indexed"},
        {"检查日期": "2026-10-14", "网址": "https://de-linker.com/sample-old-c/",
         "_coverage": "Page with redirect", "_robots": "ALLOWED", "_indexing": "INDEXING_ALLOWED",
         "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-13 17:15",
         "Google 选的正式地址": "https://de-linker.com/sample-new-c/"},
        {"检查日期": "2026-10-12", "网址": "https://de-linker.com/sample-about/",
         "_coverage": "Submitted and indexed", "_robots": "ALLOWED", "_indexing": "INDEXING_ALLOWED",
         "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-10 01:30",
         "Google 选的正式地址": "https://de-linker.com/sample-about/",
         "页面声明的正式地址": "https://de-linker.com/sample-about/",
         "备注": SAMPLE + "；站点地图里的老页面，没有执行登记"},
        {"检查日期": "2026-10-14", "网址": "https://de-linker.com/sample-about/",
         "_coverage": "Duplicate, Google chose different canonical than user", "_robots": "ALLOWED",
         "_indexing": "INDEXING_ALLOWED", "_fetch": "SUCCESSFUL", "上次抓取时间（北京时间）": "2026-10-13 08:20",
         "Google 选的正式地址": "https://de-linker.com/sample-about-us/",
         "页面声明的正式地址": "https://de-linker.com/sample-about/",
         "备注": SAMPLE + "；演示「掉出收录」：上次已收录，这次 Google 改选了另一个网址"},
    ],
    "网页表现": [
        {"日期（太平洋时间）": "2026-10-10", "网址": "https://de-linker.com/sample-about/", "_type": "web",
         "_country": "deu", "_device": "DESKTOP", "展示": 5, "点击": 0, "平均排名": 22.0, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-10", "网址": "https://de-linker.com/sample-about/", "_type": "web",
         "_country": "aut", "_device": "MOBILE", "展示": 2, "点击": 0, "平均排名": 25.5, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-11", "网址": "https://de-linker.com/sample-about/", "_type": "web",
         "_country": "deu", "_device": "DESKTOP", "展示": 3, "点击": 1, "平均排名": 18.0, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-11", "网址": "https://de-linker.com/sample-about/", "_type": "image",
         "_country": "deu", "_device": "MOBILE", "展示": 1, "点击": 0, "平均排名": 7.0, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-12", "网址": "https://de-linker.com/sample-guide-a/", "_type": "web",
         "_country": "deu", "_device": "DESKTOP", "展示": 2, "点击": 0, "平均排名": 61.5, "数据状态": "初步"},
        {"日期（太平洋时间）": "2026-10-13", "网址": "https://de-linker.com/sample-guide-a/", "_type": "web",
         "_country": "deu", "_device": "MOBILE", "展示": 3, "点击": 1, "平均排名": 47.0, "数据状态": "初步"},
        {"日期（太平洋时间）": "2026-10-13", "网址": "https://de-linker.com/sample-guide-a/", "_type": "web",
         "_country": "che", "_device": "DESKTOP", "展示": 2, "点击": 0, "平均排名": 50.0, "数据状态": "初步"},
    ],
    "搜索词": [
        {"日期（太平洋时间）": "2026-10-10", "网址": "https://de-linker.com/sample-about/", "搜索词": "beispiel suchbegriff a",
         "_type": "web", "_country": "deu", "_device": "DESKTOP", "展示": 3, "点击": 0, "平均排名": 21.0, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-11", "网址": "https://de-linker.com/sample-about/", "搜索词": "beispiel suchbegriff b",
         "_type": "web", "_country": "deu", "_device": "DESKTOP", "展示": 2, "点击": 1, "平均排名": 12.0, "数据状态": "定稿"},
        {"日期（太平洋时间）": "2026-10-12", "网址": "https://de-linker.com/sample-guide-a/", "搜索词": "beispiel suchbegriff c",
         "_type": "web", "_country": "deu", "_device": "DESKTOP", "展示": 2, "点击": 0, "平均排名": 61.5, "数据状态": "初步"},
        {"日期（太平洋时间）": "2026-10-13", "网址": "https://de-linker.com/sample-guide-a/", "搜索词": "beispiel suchbegriff c",
         "_type": "web", "_country": "deu", "_device": "MOBILE", "展示": 2, "点击": 1, "平均排名": 45.0, "数据状态": "初步"},
        {"日期（太平洋时间）": "2026-10-13", "网址": "https://de-linker.com/sample-guide-a/", "搜索词": "example search term d",
         "_type": "web", "_country": "che", "_device": "DESKTOP", "展示": 1, "点击": 0, "平均排名": 52.0, "数据状态": "初步"},
    ],
    "邮件反馈": [
        {"事件日期": "2026-10-12", "登记编号": "E-20261012-02", "事件类型": "已发送里找到"},
        {"事件日期": "2026-10-13", "登记编号": "E-20261012-02", "事件类型": "未见退信（推断已送达）"},
        {"事件日期": "2026-10-14", "登记编号": "E-20261012-02", "事件类型": "真人回复",
         "回复分类": "要报价、规格、价格或样品", "回信 Message-ID": "<sample-reply-001@example.invalid>"},
        {"事件日期": "2026-10-13", "登记编号": "E-20261013-01", "事件类型": "已发送里找到"},
        {"事件日期": "2026-10-13", "登记编号": "E-20261013-01", "事件类型": "永久退信", "退信状态码": "5.1.1",
         "回信 Message-ID": "<sample-bounce-002@example.invalid>",
         "备注": SAMPLE + "；5.1.1 = 收件邮箱不存在，这个地址停用"},
    ],
    "询盘": [
        {"询盘内部编号": 90001, "单号": "90001", "公司": "样例公司甲 GmbH", "创建时刻（北京时间）": "2026-10-14 21:12",
         "状态": "询盘已收到", "归因登记编号": "E-20261012-02",
         "归因依据": "同一公司回信要报价后，当天提交了形式发票申请"},
        {"询盘内部编号": 90002, "单号": "90002", "公司": "样例公司丙 KG", "创建时刻（北京时间）": "2026-10-11 22:30",
         "状态": "形式发票已开", "形式发票签发时刻（北京时间）": "2026-10-13 18:00",
         "备注": SAMPLE + "；找不到是哪次动作带来的，未归因，但照样算进询盘总数"},
    ],
}

LEVEL_RANK = {"L0": 0, "未到 L1": 0.5, "L1": 1, "L2": 2, "L3": 3}
EVENT_LEVEL = {
    "已发送里找到": "L0", "未见退信（推断已送达）": "L1", "暂时退信": "未到 L1", "永久退信": "未到 L1",
    "认证或信誉拒收": "未到 L1", "改址通知": "未到 L1", "自动回复": "L1", "真人回复": "L2",
}
SAMPLE_WRITE_AT = "2026-10-14 14:00"
SAMPLE_RECOMPUTE_AT = "2026-10-14 14:05"
UNKNOWN_STATE = "不在对照表里（看英文原文）"
NO_DEVICE = "不区分（Discover 不给设备）"


def prop(schema: dict, sheet: str, col: str) -> dict:
    return schema["$defs"][sheet]["properties"][col]


def latest_registration(log: list[dict], url: str, day: str) -> str:
    hits = [r for r in log if url in (r.get("网址"), r.get("旧网址")) and r["执行日期"] <= day]
    return max(hits, key=lambda r: (r["执行日期"], r["登记编号"]))["登记编号"] if hits else ""


def language(url: str) -> str:
    return "德文" if "/de/" in url else "英文"


def translate_inspection(schema: dict, raw: dict) -> dict:
    """URL Inspection result (raw API values under _keys) -> 网页收录 columns, without 和上次比."""
    r = {k: v for k, v in raw.items() if not k.startswith("_")}
    states = prop(schema, "网页收录", "收录状态")
    state = states["x-map"].get(raw["_coverage"], UNKNOWN_STATE)
    r["Google 状态原文"] = raw["_coverage"]
    r["收录状态"] = state
    r["收录结论"] = states["x-states"][state]["收录结论"]
    for col, key in (("robots 是否放行", "_robots"), ("是否允许收录", "_indexing"), ("抓取结果", "_fetch")):
        if raw.get(key):
            r[col] = prop(schema, "网页收录", col)["x-map"].get(raw[key], raw[key])
    parts = [f"上次抓取：{r['抓取结果']}" if r.get("抓取结果") else "Google 还没来抓过"]
    if r.get("robots 是否放行"):
        parts.append(f"robots：{r['robots 是否放行']}")
    if r.get("是否允许收录"):
        parts.append(f"收录指令：{r['是否允许收录']}")
    g, u = r.get("Google 选的正式地址"), r.get("页面声明的正式地址")
    if g and u and g != u:
        parts.append("Google 选的正式地址和页面声明的不同")
    r["原因说明"] = states["x-states"][state]["意思"] + "；".join(parts) + "。"
    r["反馈级别"] = "L1" if r["收录结论"] in ("已收录", "部分问题") else "L0"
    return r


def compare_with_previous(rows: list[dict]) -> None:
    """Fill 和上次比 by comparing each check with the previous check of the same URL."""
    indexed = ("已收录", "部分问题")
    last: dict[str, dict] = {}
    for r in sorted(rows, key=lambda r: (r["网址"], r["检查日期"])):
        prev = last.get(r["网址"])
        if prev is None:
            r["和上次比"] = "首次检查"
        elif (prev["收录结论"] in indexed) != (r["收录结论"] in indexed):
            r["和上次比"] = "新收录" if r["收录结论"] in indexed else "掉出收录"
        elif prev["收录状态"] != r["收录状态"]:
            r["和上次比"] = "状态变了（仍已收录）" if r["收录结论"] in indexed else "状态变了（仍未收录）"
        else:
            r["和上次比"] = "没变"
        last[r["网址"]] = r


def translate_search_row(schema: dict, sheet: str, raw: dict) -> dict:
    r = {k: v for k, v in raw.items() if not k.startswith("_")}
    r["搜索类型"] = prop(schema, sheet, "搜索类型")["x-map"][raw["_type"]]
    r["国家"] = prop(schema, sheet, "国家")["x-map"].get(raw["_country"], raw["_country"].upper())
    r["设备"] = prop(schema, sheet, "设备")["x-map"][raw["_device"]] if raw.get("_device") else NO_DEVICE
    r["点击率"] = round(r["点击"] / r["展示"], 4) if r["展示"] else None
    r["反馈级别"] = "L2" if r["点击"] > 0 else "L1"
    return r


def action_progress(data: dict[str, list[dict]], now: str) -> list[dict]:
    """Recompute 动作进度 from the other tables (the pullback scripts reuse this)."""
    perf_days = [r["日期（太平洋时间）"] for r in data["网页表现"]]
    window_start = (date.fromisoformat(max(perf_days)) - timedelta(days=6)).isoformat() if perf_days else ""
    out = []
    for e in data["执行登记"]:
        eid = e["登记编号"]
        idx = [r for r in data["网页收录"] if r.get("登记编号") == eid]
        perf = [r for r in data["网页表现"] if r.get("登记编号") == eid]
        queries = [r for r in data["搜索词"] if r.get("登记编号") == eid and r["日期（太平洋时间）"] >= window_start]
        mail = [r for r in data["邮件反馈"] if r.get("登记编号") == eid]
        inq = [r for r in data["询盘"] if r.get("归因登记编号") == eid]
        levels = ["L0"] + [r["反馈级别"] for r in idx + perf + mail] + (["L3"] if inq else [])
        own_idx = [r for r in idx if r["网址"] == e.get("网址")] or idx
        latest = max(own_idx, key=lambda r: r["检查日期"]) if own_idx else {}
        recent = [r for r in perf if r["日期（太平洋时间）"] >= window_start]
        de = [r for r in recent if r["国家"] == "德国"]
        top: dict[str, int] = {}
        for r in queries:
            top[r["搜索词"]] = top.get(r["搜索词"], 0) + r["展示"]
        dates = [r["检查日期"] for r in idx] + [r["日期（太平洋时间）"] for r in perf] \
            + [r["事件日期"] for r in mail] + [r["创建时刻（北京时间）"][:10] for r in inq]
        is_web = e["渠道"] == "网站"
        out.append({
            "登记编号": eid, "执行日期": e["执行日期"], "线": e["线"], "渠道": e["渠道"], "动作大类": e["动作大类"],
            "动作": e["动作"], "内容说明": e["内容说明"], "网址或收件公司": e.get("网址") or e.get("收件公司", ""),
            "查证结果": e["查证结果"],
            "当前最高级别": max(levels, key=LEVEL_RANK.__getitem__),
            "收录状态": latest.get("收录状态", ""), "和上次比": latest.get("和上次比", ""),
            "近 7 天展示": sum(r["展示"] for r in recent) if is_web else None,
            "近 7 天点击": sum(r["点击"] for r in recent) if is_web else None,
            "近 7 天展示（德国）": sum(r["展示"] for r in de) if is_web else None,
            "近 7 天点击（德国）": sum(r["点击"] for r in de) if is_web else None,
            "近 7 天主要搜索词": "；".join(f"{q}（{n}）" for q, n in sorted(top.items(), key=lambda x: -x[1])[:3]),
            "最新邮件事件": max(mail, key=lambda r: (r["事件日期"], LEVEL_RANK[r["反馈级别"]]))["事件类型"] if mail else "",
            "归因询盘数": len(inq),
            "最近一次反馈日期": max(dates) if dates else "",
            "更新时刻（北京时间）": now,
        })
    return out


def sample_data(schema: dict) -> dict[str, list[dict]]:
    raw = {k: [dict(r) for r in v] for k, v in SAMPLE_ROWS.items()}
    log = raw["执行登记"]
    categories = prop(schema, "执行登记", "动作")["x-category"]
    for r in log:
        r["动作大类"] = categories[r["动作"]]
    data: dict[str, list[dict]] = {"执行登记": log, "邮件反馈": raw["邮件反馈"], "询盘": raw["询盘"]}
    data["网页收录"] = [translate_inspection(schema, r) for r in raw["网页收录"]]
    compare_with_previous(data["网页收录"])
    for sheet in ("网页表现", "搜索词"):
        data[sheet] = [translate_search_row(schema, sheet, r) for r in raw[sheet]]
    for sheet, day_col in (("网页收录", "检查日期"), ("网页表现", "日期（太平洋时间）"), ("搜索词", "日期（太平洋时间）")):
        for r in data[sheet]:
            r["登记编号"] = latest_registration(log, r["网址"], r[day_col])
            r["语言"] = language(r["网址"])
    by_id = {r["登记编号"]: r for r in log}
    for r in data["邮件反馈"]:
        r["收件公司"] = by_id[r["登记编号"]].get("收件公司", "")
        r["反馈级别"] = EVENT_LEVEL[r["事件类型"]]
    for sheet in ("网页收录", "网页表现", "搜索词", "邮件反馈", "询盘"):
        for r in data[sheet]:
            r["写入时刻（北京时间）"] = SAMPLE_WRITE_AT
    for sheet in ("执行登记", "网页收录", "网页表现", "搜索词", "邮件反馈", "询盘"):
        for r in data[sheet]:
            r.setdefault("备注", SAMPLE)
    data["动作进度"] = action_progress(data, SAMPLE_RECOMPUTE_AT)
    for name in schema["x-sheet-order"]:
        unknown = {k for row in data[name] for k in row} - set(columns(schema["$defs"][name]))
        if unknown:
            raise SystemExit(f"{name}: sample columns not in schema: {sorted(unknown)}")
    return data


def data_widths(cols: list[str], rows: list[dict]) -> list[int]:
    def width(text: str) -> int:
        return sum(2 if ord(ch) > 255 else 1 for ch in text) + 2
    return [max(12, min(60, max([width(c)] + [width(str(r.get(c) or "")) for r in rows]))) for c in cols]


# ---------------------------------------------------------------- xlsx


def write_xlsx(schema: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fonts = {"title": Font(bold=True, size=16), "h2": Font(bold=True, size=12), "header": Font(bold=True)}
    wrap = Alignment(wrap_text=True, vertical="top")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def render(ws, page: Page, ws_widths):
        span = merge_span(ws_widths)
        for r, (cells, style) in enumerate(page, start=1):
            for c, v in enumerate(cells, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment = wrap
                if style in fonts:
                    cell.font = fonts[style]
                if style == "header":
                    cell.fill = PatternFill("solid", start_color=GREY)
                if style in WRITER_FILL and c == 1:
                    cell.fill = PatternFill("solid", start_color=WRITER_FILL[style])
            if len(cells) == 1:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
                ws.row_dimensions[r].height = text_row_px(cells[0], style, ws_widths) * 0.75
        widths(ws, ws_widths)

    data = sample_data(schema)
    wb = Workbook()
    home = wb.active
    home.title = "首页"
    render(home, home_page(schema, sample=True), HOME_WIDTHS)
    opt_ws = wb.create_sheet("_选项")
    opt_ws.sheet_state = "hidden"
    options: list[list[str]] = []
    for name in schema["x-sheet-order"]:
        spec = schema["$defs"][name]
        cols = columns(spec)
        ws = wb.create_sheet(name)
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=c, value=col)
            cell.font = fonts["header"]
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.fill = PatternFill("solid", start_color=WRITER_FILL[writer_key(spec["properties"][col])])
        for r, row in enumerate(data[name], start=2):
            for c, col in enumerate(cols, start=1):
                v = row.get(col)
                ws.cell(row=r, column=c, value=None if v in ("", None) else v).alignment = wrap
        for c, col in enumerate(cols, start=1):
            meta = spec["properties"][col]
            if "enum" in meta:
                letter = get_column_letter(c)
                inline = '"' + ",".join(meta["enum"]) + '"'
                if len(inline) > 250:  # Excel caps inline lists at 255 chars; long lists live on a hidden tab
                    options.append(meta["enum"])
                    col_letter = get_column_letter(len(options))
                    for i, v in enumerate(meta["enum"], start=1):
                        opt_ws.cell(row=i, column=len(options), value=v)
                    inline = f"'_选项'!${col_letter}$1:${col_letter}${len(meta['enum'])}"
                dv = DataValidation(type="list", formula1=inline, allow_blank=True)
                dv.add(f"{letter}2:{letter}{VALIDATION_ROWS}")
                ws.add_data_validation(dv)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 34
        widths(ws, data_widths(cols, data[name]))
        render(wb.create_sheet(f"{name}说明"), guide_page(name, spec), GUIDE_WIDTHS)
    wb.move_sheet(opt_ws, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("_选项"))
    wb.save(XLSX_OUT)
    print(f"wrote {XLSX_OUT.relative_to(ROOT)}")


# ---------------------------------------------------------------- Google Sheets


def rgb(hex6: str) -> dict:
    return {"red": int(hex6[0:2], 16) / 255, "green": int(hex6[2:4], 16) / 255, "blue": int(hex6[4:6], 16) / 255}


def load_target(target: str) -> str:
    path = Path(os.environ.get("PLAYBOOK_CONFIG", CONFIG_DIR / "config.toml"))
    if not path.exists():
        raise SystemExit(f"missing config {path} (see this script's docstring)")
    cfg = tomllib.loads(path.read_text(encoding="utf-8")).get("base_table", {})
    key = "spreadsheet_id" if target == "main" else "sample_spreadsheet_id"
    if not cfg.get(key):
        raise SystemExit(f"{path}: [base_table] {key} is empty")
    return cfg[key]


def google_service(api: str, version: str, scopes: list[str]):
    import httplib2
    from google.oauth2 import service_account
    from google_auth_httplib2 import AuthorizedHttp
    from googleapiclient.discovery import build

    cred = Path(os.environ.get("GOOGLE_APPLICATION_CREDENTIALS", CONFIG_DIR / "service-account.json"))
    creds = service_account.Credentials.from_service_account_file(str(cred), scopes=scopes)
    http = AuthorizedHttp(creds, http=httplib2.Http(timeout=HTTP_TIMEOUT))
    return build(api, version, http=http, cache_discovery=False)


def sheets_service():
    return google_service("sheets", "v4", SCOPES)


def beijing_now() -> datetime:
    return datetime.now(BEIJING)


def stamp(t: datetime) -> str:
    return t.astimezone(BEIJING).strftime("%Y-%m-%d %H:%M")


def row_key(spec: dict, r: dict) -> tuple:
    return tuple("" if r.get(k) is None else str(r[k]) for k in spec["x-key"])


def to_cell(v):
    return "" if v is None else v


def read_table(svc, sid: str, schema: dict, sheet: str) -> list[dict]:
    """Data rows as dicts (empty cells left out, number columns as numbers); each carries its sheet row in _row."""
    spec = schema["$defs"][sheet]
    cols = columns(spec)
    got = svc.spreadsheets().values().get(spreadsheetId=sid, range=a1(sheet, "A2:ZZ"),
                                          valueRenderOption="UNFORMATTED_VALUE").execute(num_retries=RETRIES)
    out = []
    for i, vals in enumerate(got.get("values", []), start=2):
        r: dict = {"_row": i}
        for col, v in zip(cols, vals):
            if v == "":
                continue
            kind = spec["properties"][col]["type"]
            r[col] = int(v) if kind == "integer" else float(v) if kind == "number" else str(v)
        if len(r) > 1:
            out.append(r)
    return out


_grid_rows: dict[tuple[str, str], int] = {}  # (sid, sheet) -> rowCount seen this process


def ensure_grid(svc, sid: str, sheet: str, max_row: int) -> None:
    """Grow the tab's row grid before a write would pass its end (build leaves tabs at the default 1000 rows)."""
    if max_row <= _grid_rows.get((sid, sheet), 0):
        return
    fields = "sheets(properties(sheetId,title,gridProperties(rowCount)))"
    props = next((s["properties"] for s in svc.spreadsheets().get(spreadsheetId=sid, fields=fields)
                  .execute(num_retries=RETRIES).get("sheets", []) if s["properties"]["title"] == sheet), None)
    if props is None:
        raise SystemExit(f"底表里没有「{sheet}」页")
    rows_now = props["gridProperties"]["rowCount"]
    if max_row <= rows_now:
        _grid_rows[(sid, sheet)] = rows_now
        return
    rows_new = max(rows_now * 2, max_row + 1000)
    req = {"updateSheetProperties": {"properties": {"sheetId": props["sheetId"],
                                                    "gridProperties": {"rowCount": rows_new}},
                                     "fields": "gridProperties.rowCount"}}
    svc.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": [req]}).execute(num_retries=RETRIES)
    _grid_rows[(sid, sheet)] = rows_new


def write_rows(svc, sid: str, sheet: str, data: list[tuple[int, list]]) -> None:
    """Write whole rows at fixed row numbers (never values.append: it guesses where the table is)."""
    if data:
        ensure_grid(svc, sid, sheet, max(n for n, _ in data))
    for i in range(0, len(data), 2000):
        body = {"valueInputOption": "RAW",
                "data": [{"range": a1(sheet, f"A{n}"), "values": [vals]} for n, vals in data[i:i + 2000]]}
        svc.spreadsheets().values().batchUpdate(spreadsheetId=sid, body=body).execute(num_retries=RETRIES)


def upsert_rows(svc, sid: str, schema: dict, sheet: str, rows: list[dict], now: str,
                existing: list[dict] | None = None) -> tuple[int, int]:
    """Insert new rows, overwrite rows whose key exists and whose content changed. Returns (added, updated)."""
    spec = schema["$defs"][sheet]
    cols = columns(spec)
    stamp_col = next((c for c in cols if c.startswith("写入时刻")), None)
    if existing is None:
        existing = read_table(svc, sid, schema, sheet)
    by_key = {row_key(spec, r): r for r in existing}
    next_row = max((r["_row"] for r in existing), default=1) + 1
    writes, added, updated = [], 0, 0
    for r in rows:
        old = by_key.get(row_key(spec, r))
        same = old is not None and all(_same(old.get(c), r.get(c)) for c in cols if c != stamp_col)
        if same:
            continue
        r = {**r, stamp_col: now} if stamp_col else r
        vals = [to_cell(r.get(c)) for c in cols]
        if old is None:
            writes.append((next_row, vals))
            by_key[row_key(spec, r)] = {**r, "_row": next_row}
            next_row += 1
            added += 1
        else:
            writes.append((old["_row"], vals))
            updated += 1
    write_rows(svc, sid, sheet, writes)
    return added, updated


def _same(a, b) -> bool:
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-9
    return str(a) == str(b)


def replace_rows(svc, sid: str, schema: dict, sheet: str, rows: list[dict]) -> None:
    cols = columns(schema["$defs"][sheet])
    svc.spreadsheets().values().clear(spreadsheetId=sid, range=a1(sheet, "A2:ZZ")).execute(num_retries=RETRIES)
    write_rows(svc, sid, sheet, [(i, [to_cell(r.get(c)) for c in cols]) for i, r in enumerate(rows, start=2)])


def refresh_progress(svc, sid: str, schema: dict, now: str) -> int:
    """Re-attach 登记编号 on the web tables by URL, then recompute 动作进度. Returns rows in 动作进度."""
    data = {name: read_table(svc, sid, schema, name) for name in schema["x-sheet-order"] if name != "动作进度"}
    log = data["执行登记"]
    for sheet, day_col in (("网页收录", "检查日期"), ("网页表现", "日期（太平洋时间）"), ("搜索词", "日期（太平洋时间）")):
        before, moved = [dict(r) for r in data[sheet]], []
        for r in data[sheet]:
            eid = latest_registration(log, r["网址"], r[day_col]) or None
            if eid != r.get("登记编号"):
                r["登记编号"] = eid
                moved.append(r)
        if moved:
            upsert_rows(svc, sid, schema, sheet, moved, now, existing=before)
    progress = action_progress(data, now)
    replace_rows(svc, sid, schema, "动作进度", progress)
    return len(progress)


def a1(title: str, rng: str = "A1") -> str:
    return "'" + title.replace("'", "''") + "'!" + rng


def tab_order(schema: dict) -> list[str]:
    order = ["首页"]
    for name in schema["x-sheet-order"]:
        order += [name, f"{name}说明"]
    return order


def fetch_meta(svc, sid: str) -> dict[str, dict]:
    meta = svc.spreadsheets().get(
        spreadsheetId=sid,
        fields="sheets(properties(sheetId,title,index,gridProperties))",
    ).execute(num_retries=RETRIES)
    return {s["properties"]["title"]: s for s in meta["sheets"]}


def col_widths_req(sheet_id: int, widths: list[int]) -> list[dict]:
    return [{"updateDimensionProperties": {
        "range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
        "properties": {"pixelSize": px(w)}, "fields": "pixelSize"}} for i, w in enumerate(widths)]


def page_requests(sheet_id: int, page: Page, widths: list[int]) -> list[dict]:
    span = merge_span(widths)
    reqs: list[dict] = [
        {"unmergeCells": {"range": {"sheetId": sheet_id}}},
        {"updateCells": {"range": {"sheetId": sheet_id}, "fields": "userEnteredValue,userEnteredFormat"}},
        {"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": 0,
                                                 "endIndex": len(page) + 50},
                                       "properties": {"pixelSize": 21}, "fields": "pixelSize"}},
    ]
    for r, (cells, style) in enumerate(page):
        if not cells:
            continue
        row = []
        for c, v in enumerate(cells):
            fmt: dict = {"verticalAlignment": "TOP", "wrapStrategy": "WRAP"}
            if style in ("title", "h2", "header"):
                fmt["textFormat"] = {"bold": True, "fontSize": {"title": 16, "h2": 12}.get(style, 10)}
            if style == "header":
                fmt["backgroundColor"] = rgb(GREY)
            if style in WRITER_FILL and c == 0:
                fmt["backgroundColor"] = rgb(WRITER_FILL[style])
            row.append({"userEnteredValue": {"stringValue": v}, "userEnteredFormat": fmt})
        reqs.append({"updateCells": {"start": {"sheetId": sheet_id, "rowIndex": r, "columnIndex": 0},
                                     "rows": [{"values": row}], "fields": "userEnteredValue,userEnteredFormat"}})
        if len(cells) == 1:
            reqs.append({"mergeCells": {"mergeType": "MERGE_ALL", "range": {
                "sheetId": sheet_id, "startRowIndex": r, "endRowIndex": r + 1, "startColumnIndex": 0,
                "endColumnIndex": span}}})
            reqs.append({"updateDimensionProperties": {
                "range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": r, "endIndex": r + 1},
                "properties": {"pixelSize": text_row_px(cells[0], style, widths)}, "fields": "pixelSize"}})
        else:
            reqs.append({"autoResizeDimensions": {"dimensions": {
                "sheetId": sheet_id, "dimension": "ROWS", "startIndex": r, "endIndex": r + 1}}})
    return reqs


def data_sheet_requests(sheet_id: int, spec: dict, widths: list[int]) -> list[dict]:
    cols = columns(spec)
    reqs: list[dict] = [{"updateSheetProperties": {
        "properties": {"sheetId": sheet_id, "gridProperties": {"frozenRowCount": 1, "frozenColumnCount": 1}},
        "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}}]
    header = []
    for col in cols:
        header.append({"userEnteredValue": {"stringValue": col}, "userEnteredFormat": {
            "textFormat": {"bold": True}, "wrapStrategy": "WRAP", "verticalAlignment": "MIDDLE",
            "backgroundColor": rgb(WRITER_FILL[writer_key(spec["properties"][col])])}})
    reqs.append({"updateCells": {"start": {"sheetId": sheet_id, "rowIndex": 0, "columnIndex": 0},
                                 "rows": [{"values": header}], "fields": "userEnteredValue,userEnteredFormat"}})
    # all data cells as plain text, except number columns, so dates and IDs never get reinterpreted
    for c, col in enumerate(cols):
        kind = spec["properties"][col]["type"]
        fmt = {"numberFormat": {"type": "NUMBER", "pattern": "0.##"}} if kind in ("integer", "number") \
            else {"numberFormat": {"type": "TEXT"}}
        fmt.update({"wrapStrategy": "WRAP", "verticalAlignment": "TOP"})
        reqs.append({"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": VALIDATION_ROWS,
                                              "startColumnIndex": c, "endColumnIndex": c + 1},
                                    "cell": {"userEnteredFormat": fmt},
                                    "fields": "userEnteredFormat(numberFormat,wrapStrategy,verticalAlignment)"}})
        meta = spec["properties"][col]
        rule = None
        if "enum" in meta:
            rule = {"condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": v} for v in meta["enum"]]},
                    "strict": True, "showCustomUi": True}
        reqs.append({"setDataValidation": {"range": {"sheetId": sheet_id, "startRowIndex": 1,
                                                     "endRowIndex": VALIDATION_ROWS,
                                                     "startColumnIndex": c, "endColumnIndex": c + 1},
                                           **({"rule": rule} if rule else {})}})
    reqs.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "ROWS",
                                                         "startIndex": 0, "endIndex": 1},
                                               "properties": {"pixelSize": 42}, "fields": "pixelSize"}})
    return reqs + col_widths_req(sheet_id, widths)


def build(schema: dict, target: str) -> None:
    sid = load_target(target)
    svc = sheets_service()
    order = tab_order(schema)
    tabs = fetch_meta(svc, sid)

    # 1. add missing tabs; drop the untouched default tab
    reqs = [{"addSheet": {"properties": {"title": t}}} for t in order if t not in tabs]
    if reqs:
        svc.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": reqs}).execute(num_retries=RETRIES)
        tabs = fetch_meta(svc, sid)
    extra = [t for t in tabs if t not in order]
    for t in extra:
        vals = svc.spreadsheets().values().get(spreadsheetId=sid, range=a1(t, "A1:Z5")).execute(num_retries=RETRIES).get("values")
        if not vals:
            svc.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": [
                {"deleteSheet": {"sheetId": tabs[t]["properties"]["sheetId"]}}]}).execute(num_retries=RETRIES)
            print(f"removed empty tab {t}")
        else:
            print(f"kept unknown tab with data: {t}", file=sys.stderr)
    tabs = fetch_meta(svc, sid)

    # 2. data tabs: refuse to touch a header that has data under it and differs from the schema
    data = sample_data(schema) if target == "sample" else None
    for name in schema["x-sheet-order"]:
        cols = columns(schema["$defs"][name])
        got = svc.spreadsheets().values().get(spreadsheetId=sid, range=a1(name, "A1:ZZ2")).execute(num_retries=RETRIES).get("values", [])
        has_rows = len(got) > 1 and any(got[1])
        if has_rows and got[0] != cols and target == "main":
            raise SystemExit(f"{name}: header differs from schema and the tab already has data; migrate by hand")
        if got and got[0] != cols:  # drop stale header cells beyond the new column count
            svc.spreadsheets().values().clear(spreadsheetId=sid, range=a1(name, "A1:ZZ1")).execute(num_retries=RETRIES)

    title = "底表" if target == "main" else "底表（样例）"
    reqs = [{"updateSpreadsheetProperties": {"properties": {"title": title}, "fields": "title"}}]
    for i, t in enumerate(order):
        reqs.append({"updateSheetProperties": {"properties": {"sheetId": tabs[t]["properties"]["sheetId"], "index": i},
                                               "fields": "index"}})
    reqs += page_requests(tabs["首页"]["properties"]["sheetId"], home_page(schema, sample=target == "sample"), HOME_WIDTHS)
    reqs += col_widths_req(tabs["首页"]["properties"]["sheetId"], HOME_WIDTHS)
    for name in schema["x-sheet-order"]:
        spec = schema["$defs"][name]
        rows = data[name] if data else []
        reqs += data_sheet_requests(tabs[name]["properties"]["sheetId"], spec, data_widths(columns(spec), rows))
        gid = tabs[f"{name}说明"]["properties"]["sheetId"]
        reqs += page_requests(gid, guide_page(name, spec), GUIDE_WIDTHS) + col_widths_req(gid, GUIDE_WIDTHS)
    svc.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": reqs}).execute(num_retries=RETRIES)

    if data:
        for name in schema["x-sheet-order"]:
            cols = columns(schema["$defs"][name])
            svc.spreadsheets().values().clear(spreadsheetId=sid, range=a1(name, "A2:ZZ")).execute(num_retries=RETRIES)
            values = [["" if r.get(c) is None else r.get(c) for c in cols] for r in data[name]]
            if values:
                svc.spreadsheets().values().update(spreadsheetId=sid, range=a1(name, "A2"),
                                                   valueInputOption="RAW", body={"values": values}).execute(num_retries=RETRIES)
    print(f"built {target}: {len(order)} tabs")


def check(schema: dict, target: str) -> int:
    sid = load_target(target)
    svc = sheets_service()
    tabs = fetch_meta(svc, sid)
    problems = []
    order = tab_order(schema)
    missing = [t for t in order if t not in tabs]
    problems += [f"missing tab: {t}" for t in missing]
    for name in schema["x-sheet-order"]:
        if name not in tabs:
            continue
        spec = schema["$defs"][name]
        cols = columns(spec)
        got = svc.spreadsheets().values().get(spreadsheetId=sid, range=a1(name, "A1:ZZ1")).execute(num_retries=RETRIES).get("values", [[]])
        if got[0] != cols:
            problems.append(f"{name}: header differs from schema")
        grid = tabs[name]["properties"].get("gridProperties", {})
        if grid.get("frozenRowCount") != 1:
            problems.append(f"{name}: header row not frozen")
        enum_cols = {i for i, c in enumerate(cols) if "enum" in spec["properties"][c]}
        row2 = svc.spreadsheets().get(spreadsheetId=sid, ranges=[a1(name, "A2:ZZ2")], includeGridData=True,
                                      fields="sheets(data(rowData(values(dataValidation))))").execute(num_retries=RETRIES)
        cells = (row2["sheets"][0].get("data", [{}])[0].get("rowData") or [{}])[0].get("values", [])
        validated = {i for i, v in enumerate(cells)
                     if v.get("dataValidation", {}).get("condition", {}).get("type") == "ONE_OF_LIST"}
        if enum_cols - validated:
            problems.append(f"{name}: dropdown missing on {sorted(cols[i] for i in enum_cols - validated)}")
    for p in problems:
        print("FAIL", p)
    print("OK structure matches schema" if not problems else f"{len(problems)} problem(s)")
    return 1 if problems else 0


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["xlsx", "build", "check"])
    ap.add_argument("--target", choices=["main", "sample"], default="main")
    args = ap.parse_args()
    schema = load_schema()
    if args.command == "xlsx":
        write_xlsx(schema)
    elif args.command == "build":
        build(schema, args.target)
    else:
        sys.exit(check(schema, args.target))


if __name__ == "__main__":
    main()
