# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "openpyxl>=3.1",
# ]
# ///
"""
激活名单：把操作者交来的原表转成统一格式，报调研进度。名单只在私有 vault；列定义在 schemas/activation-list.schema.json。
怎么用名单（调研、联系人）在 skills/lead-company-research/references/list-mode.md；本脚本只做机械的部分。

  uv run scripts/activation_list.py convert 原件.xlsx [--out CSV]
                       读原件第一张工作表，一行一位联系人，按 company_name 原文编 list_id，写 CSV（UTF-8 带 BOM）。
                       原件不动。已有 CSV 时沿用其中公司的 list_id，新公司接着编。原件有未定义的列 → 报错退出 2。
  uv run scripts/activation_list.py progress
                       已调研 n / 共 N（n = 02-companies/activation/ 里 status: approved 的文件数），并给出下一家。
  uv run scripts/activation_list.py show --id L-0001
                       打印这家公司在名单里的全部行（调研时用）。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCHEMA = ROOT / "schemas" / "activation-list.schema.json"
LIST_DIR = ROOT / "deliverables" / "lead" / "00-system" / "activation-list"
CSV_PATH = LIST_DIR / "activation-list.csv"
RESEARCH_DIR = ROOT / "deliverables" / "lead" / "02-companies" / "activation"

# 原件表头 → 列名；列名本身也认
HEADER_ALIASES = {
    "Priority": "priority",
    "Firmenname": "company_name",
    "Vorname": "first_name",
    "Nachname": "last_name",
    "PLZ": "postal_code",
    "Adresse": "address",
    "Stadt": "city",
    "Land": "country",
    "Email": "email",
    "Telefon": "phone",
    "Homepage": "homepage",
}


def columns() -> list[str]:
    return list(json.loads(SCHEMA.read_text(encoding="utf-8"))["properties"])


def cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def company_key(row: dict) -> str:
    """同名即同一家；没写公司名的行按它在原件里的位置各算一家。"""
    return row["company_name"] or f"#{row['source_file']}:{row['source_sheet']}:{row['source_row']}"


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def convert(original: Path, out: Path) -> int:
    from openpyxl import load_workbook

    cols = columns()
    wb = load_workbook(original, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [cell(h) for h in next(rows)]
    mapped = []
    for h in header:
        name = HEADER_ALIASES.get(h, h)
        if h and name not in cols:
            print(f"原件有未定义的列「{h}」：先在 schemas/activation-list.schema.json 加列，再转。", file=sys.stderr)
            return 2
        mapped.append(name if h else None)

    ids: dict[str, str] = {}
    next_n = 1
    if out.exists():
        for r in read_csv(out):
            ids.setdefault(company_key(r), r["list_id"])
            next_n = max(next_n, int(r["list_id"][2:]) + 1)

    written = []
    for i, raw in enumerate(rows, start=2):
        values = [cell(v) for v in raw]
        if not any(values):
            continue
        row = {c: "" for c in cols}
        for name, v in zip(mapped, values):
            if name:
                row[name] = v
        row.update(source_file=original.name, source_sheet=ws.title, source_row=str(i))
        key = company_key(row)
        if key not in ids:
            ids[key] = f"L-{next_n:04d}"
            next_n += 1
        row["list_id"] = ids[key]
        written.append(row)

    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(written)
    companies = len({r["list_id"] for r in written})
    print(f"写入 {out}：{len(written)} 行（联系人），{companies} 家公司（list_id）。原件未改：{original}")
    return 0


def companies(rows: list[dict]) -> list[dict]:
    """一家一条：list_id、priority（取这家各行里最高的）、首行。按调研顺序排。"""
    by_id: dict[str, dict] = {}
    for r in rows:
        c = by_id.setdefault(r["list_id"], {"list_id": r["list_id"], "priority": "", "first": r})
        if r["priority"] and (not c["priority"] or r["priority"] < c["priority"]):
            c["priority"] = r["priority"]
    return sorted(by_id.values(), key=lambda c: (c["priority"] or "9", int(c["list_id"][2:])))


def research_status() -> dict[str, str]:
    status = {}
    for path in sorted(RESEARCH_DIR.glob("L-*.md")) if RESEARCH_DIR.exists() else []:
        text = path.read_text(encoding="utf-8")
        front = text.split("---")[1] if text.startswith("---") else ""
        m = re.search(r"^status:\s*(\S+)", front, re.M)
        status[path.stem] = m.group(1) if m else "draft"
    return status


def progress() -> int:
    if not CSV_PATH.exists():
        print("还没有激活名单（deliverables/lead/00-system/activation-list/activation-list.csv）。先 convert。")
        return 1
    comps = companies(read_csv(CSV_PATH))
    status = research_status()
    done = sum(1 for c in comps if status.get(c["list_id"]) == "approved")
    drafts = [c["list_id"] for c in comps if status.get(c["list_id"]) not in (None, "approved")]
    todo = [c for c in comps if c["list_id"] not in status]
    print(f"已调研 {done} / 共 {len(comps)}")
    if drafts:
        print(f"调研未过闸（draft/stale）：{', '.join(drafts)}")
    if todo:
        print(f"下一家：{todo[0]['list_id']}（priority {todo[0]['priority'] or '空'}）")
    return 0


def show(list_id: str) -> int:
    rows = [r for r in read_csv(CSV_PATH) if r["list_id"] == list_id]
    if not rows:
        print(f"名单里没有 {list_id}", file=sys.stderr)
        return 1
    for r in rows:
        print(json.dumps(r, ensure_ascii=False))
    return 0


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="激活名单：转格式、报调研进度")
    sub = p.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("convert")
    c.add_argument("original", type=Path)
    c.add_argument("--out", type=Path, default=CSV_PATH)
    sub.add_parser("progress")
    s = sub.add_parser("show")
    s.add_argument("--id", required=True)
    a = p.parse_args(argv)
    if a.cmd == "convert":
        return convert(a.original, a.out)
    if a.cmd == "progress":
        return progress()
    return show(a.id)


if __name__ == "__main__":
    raise SystemExit(main())
