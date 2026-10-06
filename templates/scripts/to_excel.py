"""
Lead generation Excel export utility.
Uses openpyxl to create .xlsx files with proper formatting.
Paired with header-only workbooks in templates/.
"""
from __future__ import annotations

from pathlib import Path


def create_excel(
    output_path: str | Path,
    sheets: dict[str, dict],
    metadata: dict[str, str] | None = None,
) -> Path:
    """
    Create an Excel file with formatted data.

    Args:
        output_path: Path for the .xlsx file
        sheets: Dict of sheet_name -> {"columns": [...], "rows": [[...], ...]}
        metadata: Optional dict of metadata key-value pairs for a "Metadata" sheet

    Returns:
        Path to created file
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError as exc:
        raise ImportError("openpyxl required: pip install openpyxl") from exc

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    wb = Workbook()
    wb.remove(wb.active)

    if metadata:
        ws = wb.create_sheet("Metadata")
        ws.append(["Key", "Value"])
        for key, value in metadata.items():
            ws.append([key, str(value)])
        _style_header(ws, header_font_white, header_fill, thin_border)

    for sheet_name, data in sheets.items():
        ws = wb.create_sheet(sheet_name)
        columns = data["columns"]
        rows = data.get("rows", [])

        ws.append(columns)
        for row in rows:
            ws.append(row)

        _style_header(ws, header_font_white, header_fill, thin_border)
        _auto_width(ws)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path


def _style_header(ws, header_font_white, header_fill, thin_border):
    from openpyxl.styles import Alignment

    for cell in ws[1]:
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border


def _auto_width(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)
