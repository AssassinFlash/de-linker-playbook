#!/usr/bin/env python3
"""Validate playbook deliverable contracts (schemas, JSON refs, workbooks)."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional at runtime
    load_workbook = None


SOURCE_KEYS = {"source", "evidence"}
SOURCE_HEADER_PATTERN = re.compile(r"\b(source|evidence)\b", re.IGNORECASE)
JSON_REFERENCE_PATTERN = re.compile(r"(?<![\w/-])([\w./{}-]+\.json)\b")
JSON_SCHEMA_MAP = {
    "ai-visibility-report.json": "rowset-deliverable.schema.json",
    "analytics-and-optimization.json": "rowset-deliverable.schema.json",
    "competitor-profiles.json": "rowset-deliverable.schema.json",
    "content-strategy.json": "rowset-deliverable.schema.json",
    "copy-library.json": "rowset-deliverable.schema.json",
    "cro-experiment-queue.json": "rowset-deliverable.schema.json",
    "funnel-redesign.json": "rowset-deliverable.schema.json",
    "google-ads-structure.json": "google-ads-structure.schema.json",
    "inquiry-flow-data.json": "rowset-deliverable.schema.json",
    "launch-strategy.json": "rowset-deliverable.schema.json",
    "orchestrator-run-{timestamp}.json": "orchestrator-run.schema.json",
    "orchestrator-run-{{timestamp}}.json": "orchestrator-run.schema.json",
    "paid-advertising-blueprint.json": "rowset-deliverable.schema.json",
    "paid-advertising.json": "rowset-deliverable.schema.json",
    "positioning-data.json": "rowset-deliverable.schema.json",
    "profile-eval.json": "profile-eval.schema.json",
    "profile-evidence.json": "profile-evidence.schema.json",
    "prospect-research.json": "rowset-deliverable.schema.json",
    "referral-program-design.json": "rowset-deliverable.schema.json",
    "retention-strategy.json": "rowset-deliverable.schema.json",
    "seo-audit-data.json": "rowset-deliverable.schema.json",
    "sequence-builder.json": "rowset-deliverable.schema.json",
    "site-schema-blueprint.json": "site-schema-blueprint.schema.json",
    "social-calendar.json": "rowset-deliverable.schema.json",
}
JSON_REFERENCE_EXEMPTIONS = {
    "brand-profile.json",
    "history.json",
    "eval-history.json",
    "marketing-teach.json",
}
SOURCE_EXEMPT_SHEETS = {"Metadata", "Instructions", "README", "Cover"}


def default_project_root() -> Path:
    return Path(__file__).resolve().parent.parent


def validate_schema_files(project_root: Path | None = None) -> list[str]:
    root = project_root or default_project_root()
    schema_dir = root / "schemas"
    errors = []
    for schema_name in sorted(set(JSON_SCHEMA_MAP.values())):
        schema_path = schema_dir / schema_name
        if not schema_path.exists():
            errors.append(f"Missing JSON schema: schemas/{schema_name}")
            continue
        try:
            data = json.loads(schema_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            errors.append(
                f"Invalid JSON schema {schema_path}: line {exc.lineno} "
                f"column {exc.colno}: {exc.msg}"
            )
            continue
        if data.get("$schema") is None or data.get("type") is None:
            errors.append(f"Schema {schema_path} must include $schema and type")
    return errors


def json_reference_schema_errors(project_root: Path | None = None) -> list[str]:
    root = project_root or default_project_root()
    scan_roots = [root / "skills"]
    errors = []
    seen: set[tuple[str, str]] = set()
    for scan_root in scan_roots:
        if not scan_root.exists():
            continue
        for path in sorted(scan_root.rglob("*.md")):
            text = path.read_text(encoding="utf-8")
            for match in JSON_REFERENCE_PATTERN.finditer(text):
                basename = Path(match.group(1).strip("`),.;")).name
                if basename in JSON_REFERENCE_EXEMPTIONS:
                    continue
                if basename.endswith(".schema.json"):
                    continue
                key = (basename, str(path))
                if key in seen:
                    continue
                seen.add(key)
                schema_name = JSON_SCHEMA_MAP.get(basename)
                if schema_name is None:
                    errors.append(
                        f"{path} references JSON deliverable without schema mapping: {basename}"
                    )
    return errors


def validate_workbook(path: Path) -> list[str]:
    if load_workbook is None:
        return []
    errors = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        if sheet.title in SOURCE_EXEMPT_SHEETS:
            continue
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        source_indexes = [
            index for index, header in enumerate(headers)
            if SOURCE_HEADER_PATTERN.search(header)
        ]
        data_rows = list(rows)
        non_empty_rows = [
            row for row in data_rows
            if any(value not in (None, "") for value in row)
        ]
        if not non_empty_rows:
            continue
        if not source_indexes:
            errors.append(f"{path}:{sheet.title} lacks source/evidence header")
            continue
        for offset, row in enumerate(non_empty_rows, start=2):
            if all(
                (row[index] if index < len(row) else None) in (None, "")
                for index in source_indexes
            ):
                errors.append(f"{path}:{sheet.title} row {offset} has empty source/evidence")
    workbook.close()
    return errors


def validate_deliverable_workbooks(project_root: Path | None = None) -> list[str]:
    """Check filled workbooks under deliverables/ for source/evidence columns."""
    root = project_root or default_project_root()
    deliverables = root / "deliverables"
    if not deliverables.exists():
        return []
    if load_workbook is None:
        return []
    errors = []
    for path in sorted(deliverables.rglob("*.xlsx")):
        errors.extend(validate_workbook(path))
    return errors


def validate_json_file(path: Path, schema: dict[str, Any]) -> list[str]:
    errors = []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        return [f"{path} is invalid JSON: line {exc.lineno} column {exc.colno}: {exc.msg}"]
    required = schema.get("required", [])
    if isinstance(data, dict):
        missing = [key for key in required if key not in data]
        for key in missing:
            errors.append(f"{path} missing required key: {key}")
    return errors


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv if argv is None else argv)
    if "-h" in argv or "--help" in argv:
        print("Usage: python evals/validate_deliverables.py [project_root]")
        return 0
    positional = [a for a in argv[1:] if not a.startswith("-")]
    project_root = Path(positional[0]).resolve() if positional else default_project_root()

    errors = []
    errors.extend(validate_schema_files(project_root))
    errors.extend(json_reference_schema_errors(project_root))
    errors.extend(validate_deliverable_workbooks(project_root))

    if errors:
        print("Deliverable validation failed")
        for error in errors:
            print(f"  - {error}")
        return 1

    print("Deliverable quality gates passed")
    print("  - JSON schema files are parseable")
    print("  - Promised JSON deliverables in skills/ have schema mappings")
    print("  - Filled workbooks under deliverables/ include source/evidence")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
